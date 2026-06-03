import pandas as pd
import matplotlib.pyplot as plt
import os
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from mlxtend.frequent_patterns import apriori, association_rules

# Setting default matplotlib style
plt.style.use('default')

# ==========================================
# 1. LOAD & PREPARE DATA (DARI EXCEL)
# ==========================================
print("1. Membaca data transaksi dari file Excel...")
file_path = 'data_penjualan.xlsx' # Menggunakan file Excel

try:
    df = pd.read_excel(file_path, sheet_name='Transaksi')
    print("✅ Dataset Excel berhasil dimuat!")
except FileNotFoundError:
    print(f"❌ Error: File '{file_path}' tidak ditemukan. Pastikan file ada di dalam folder proyek.")
    raise
except Exception as e:
    print(f"❌ Terjadi kesalahan saat memuat file: {e}")
    raise

df['Tanggal'] = pd.to_datetime(df['tgl_transaksi'])

# Agregasi harian awal
daily_df = df.groupby(['tgl_transaksi', 'kode_produk', 'nama_produk'])['total_nilai'].sum().reset_index()
daily_df = daily_df.sort_values(by=['kode_produk', 'tgl_transaksi'])

# ==========================================
# 2. PERHITUNGAN MA & IDENTIFIKASI TREN
# ==========================================
print("2. Menghitung Moving Average dan Tren Rising Star (Metode Vektorisasi)...")
window = 3
daily_df['MA'] = daily_df.groupby('kode_produk')['total_nilai'].transform(lambda x: x.rolling(window=window).mean())

# Tentukan tren naik (Is_Rising)
daily_df['Is_Rising'] = daily_df.groupby('kode_produk')['MA'].diff() > 0
daily_df['Trend_Session'] = (daily_df['Is_Rising'] != daily_df.groupby('kode_produk')['Is_Rising'].shift()).groupby(daily_df['kode_produk']).cumsum()

# Hitung durasi kenaikan berurutan TANPA .apply() untuk mencegah hilangnya kolom 'kode_produk'
daily_df['Consecutive_Rise'] = daily_df.groupby(['kode_produk', 'Trend_Session']).cumcount() + 1
daily_df.loc[daily_df['Is_Rising'] == False, 'Consecutive_Rise'] = 0

# ==========================================
# 3. NORMALISASI (BASE 100)
# ==========================================
# Mencari nilai MA valid pertama per produk dan melakukan normalisasi TANPA .apply()
first_valid_ma = daily_df.dropna(subset=['MA']).groupby('kode_produk')['MA'].first()
daily_df['First_MA'] = daily_df['kode_produk'].map(first_valid_ma).fillna(1)

daily_df['Normalized'] = (daily_df['MA'] / daily_df['First_MA']) * 100

# Bersihkan kolom bantuan
daily_df = daily_df.drop(columns=['First_MA'])

# ==========================================
# 4. PERHITUNGAN GROWTH % & FILTER 12 HARI
# ==========================================
rising_sessions = daily_df[daily_df['Is_Rising'] == True].copy()

growth_per_session = rising_sessions.groupby(['kode_produk', 'nama_produk', 'Trend_Session']).agg(
    Growth_Pct=('MA', lambda x: (x.iloc[-1] / x.iloc[0] - 1) * 100),
    Max_Consecutive=('Consecutive_Rise', 'max')
).reset_index()

# Filter minimal tren 12 hari berturut-turut
target_rentetan = 12
mask_12_days = growth_per_session['Max_Consecutive'] >= target_rentetan
filtered_growth = growth_per_session[mask_12_days].copy()

# Ambil growth tertinggi jika ada produk yang tren >12 harinya terjadi lebih dari sekali
final_growth = filtered_growth.groupby(['kode_produk', 'nama_produk'])['Growth_Pct'].max().reset_index()

# Merge dengan total penjualan riil
total_sales = df.groupby('kode_produk')['total_nilai'].sum().reset_index()
final_report = pd.merge(final_growth, total_sales, on='kode_produk')
final_report = final_report.sort_values(by='Growth_Pct', ascending=False)

# Siapkan dataframe untuk di-export ke Excel
display_final = final_report.copy()
display_final.columns = ['Kode Produk', 'Nama Produk', 'Growth %', 'Total Penjualan']
display_final['Growth %'] = display_final['Growth %'].round(2)
display_final['Total Penjualan'] = display_final['Total Penjualan'].astype(int)

# ==========================================
# 5. APRIORI / POTENTIAL PACKAGING
# ==========================================
print("3. Memproses Apriori Analysis untuk Potential Packaging...")

basket = df.groupby(['nomor_struk', 'nama_produk'])['jumlah_terjual'].sum().unstack(fill_value=0)
basket = (basket > 0).astype(int)

frequent_itemsets = apriori(basket, min_support=0.01, use_colnames=True)
rules = association_rules(frequent_itemsets, metric='lift', min_threshold=1)

# Filter 1: Harus mengandung Rising Star
rising_star_products = set(final_report['nama_produk'])

def contains_rising_star(itemset):
    return any(item in rising_star_products for item in itemset)

rules = rules[
    rules['antecedents'].apply(contains_rising_star) |
    rules['consequents'].apply(contains_rising_star)
].copy()

# Filter 2: Lift minimal 2 & Sorting
rules = rules[rules['lift'] >= 2]
rules = rules.sort_values(by=['lift', 'support', 'confidence'], ascending=[False, False, False])

# Hitung jumlah invoice transaksi per rule
total_transactions = df['nomor_struk'].nunique()
rules['Jumlah_Transaksi_Rule'] = (rules['support'] * total_transactions).round(0).astype(int)

packaging_result = pd.DataFrame({
    'Jika Membeli': rules['antecedents'].apply(lambda x: ', '.join(sorted(x, reverse=True))),
    'Maka Membeli': rules['consequents'].apply(lambda x: ', '.join(sorted(x, reverse=True))),
    'Jumlah Invoice': rules['Jumlah_Transaksi_Rule'],
    'Support': rules['support'].round(2),
    'Confidence': rules['confidence'].round(2),
    'Lift': rules['lift'].round(2)
})

# ==========================================
# 6. EXPORT KE EXCEL (retail_insight.xlsx)
# ==========================================
print("4. Mengekspor hasil analisis ke file 'retail_insight.xlsx'...")
output_file = 'retail_insight.xlsx' 

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    display_final.to_excel(writer, sheet_name='Rising Star', index=False)
    packaging_result.to_excel(writer, sheet_name='Potential Packaging', index=False)

# Styling Excel agar rapi
workbook = load_workbook(output_file)
for sheet_name in ['Rising Star', 'Potential Packaging']:
    worksheet = workbook[sheet_name]
    
    # Header tebal
    for cell in worksheet[1]: 
        cell.font = Font(bold=True)
    
    # Auto-adjust lebar kolom
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        worksheet.column_dimensions[column_letter].width = max_length + 3

# Format angka ribuan untuk "Total Penjualan"
worksheet = workbook['Rising Star']
for row in worksheet.iter_rows(min_row=2):
    row[3].number_format = '#,##0' 

workbook.save(output_file)

# ==========================================
# 7. VISUALISASI MATPLOTLIB
# ==========================================
print("5. Menghasilkan visualisasi Matplotlib...")

plot_df = daily_df[daily_df['kode_produk'].isin(final_report['kode_produk'])].copy()

if not plot_df.empty:
    custom_palette = ['#FFD700', '#2ecc71', '#3498db', '#9b59b6', '#e74c3c'] 
    
    color_mapping = {}
    rank_mapping = {}
    for i, row in enumerate(final_report.itertuples()):
        color_mapping[row.kode_produk] = custom_palette[i] if i < len(custom_palette) else '#34495e'
        rank_mapping[row.kode_produk] = i + 1

    top3_sales = df.groupby(['kode_produk', 'nama_produk'])['total_nilai'].sum().reset_index().sort_values(by='total_nilai', ascending=False).head(3)
    top3_codes = top3_sales['kode_produk'].tolist()
    top3_plot_df = daily_df[daily_df['kode_produk'].isin(top3_codes)].copy()
    grey_colors = ['#B0B0B0', '#909090', '#707070']

    # --- GRAFIK 1: PERTUMBUHAN RELATIF (BASE 100) ---
    fig, ax = plt.subplots(figsize=(15, 8), dpi=100)
    
    # Plot baseline Top 3 Sales
    for idx, (kode_produk, group) in enumerate(top3_plot_df.groupby('kode_produk')):
        ax.plot(group['tgl_transaksi'], group['Normalized'], linestyle='--', linewidth=2, marker='o', markersize=3, color=grey_colors[idx], alpha=0.7, label=f"Top Sales: {group['nama_produk'].iloc[0]}")
    
    # Plot target Rising Star
    for kode_produk, group in plot_df.groupby('kode_produk'):
        rank = rank_mapping.get(kode_produk)
        ax.plot(group['tgl_transaksi'], group['Normalized'], marker='o', markersize=4, linewidth=2.5, color=color_mapping.get(kode_produk), label=f"Rank {rank}: {group['nama_produk'].iloc[0]}")

    ax.set_title('ANALISIS PERTUMBUHAN RELATIF PRODUK RISING STAR\n(Dengan Benchmark Top 3 Total Penjualan)', fontweight='bold', fontsize=16, pad=20)
    ax.set_xlabel('Periode Tanggal', fontsize=12, labelpad=10)
    ax.set_ylabel('Indeks Pertumbuhan (Base 100)', fontsize=12, labelpad=10)
    ax.grid(True, linestyle='--', linewidth=0.5, alpha=0.5)
    ax.axhline(y=100, color='black', linestyle='-', linewidth=1, alpha=0.5)
    plt.xticks(rotation=45, ha='right')
    
    # Susun ulang legend (Top Sales dulu, lalu Rank Rising Star)
    handles, labels = ax.get_legend_handles_labels()
    top_sales_leg = [x for x in zip(handles, labels) if 'Top Sales' in x[1]]
    rising_leg = sorted([x for x in zip(handles, labels) if 'Rank' in x[1]], key=lambda x: int(x[1].split(':')[0].split()[1]))
    final_handles, final_labels = zip(*(top_sales_leg + rising_leg))
    ax.legend(final_handles, final_labels, title="Kategori Produk", bbox_to_anchor=(1.02, 1), loc='upper left')

    plt.tight_layout()
    plt.savefig('rising_star_index.png', bbox_inches='tight')

    # --- GRAFIK 2: NILAI PENJUALAN ASLI ---
    fig2, ax2 = plt.subplots(figsize=(15, 8), dpi=100)
    
    for idx, (kode_produk, group) in enumerate(top3_plot_df.groupby('kode_produk')):
        ax2.plot(group['tgl_transaksi'], group['total_nilai'], linestyle='--', linewidth=2, marker='o', markersize=3, color=grey_colors[idx], alpha=0.7, label=f"Top Sales: {group['nama_produk'].iloc[0]}")
    
    for kode_produk, group in plot_df.groupby('kode_produk'):
        rank = rank_mapping.get(kode_produk)
        ax2.plot(group['tgl_transaksi'], group['total_nilai'], marker='o', markersize=4, linewidth=2.5, color=color_mapping.get(kode_produk), label=f"Rank {rank}: {group['nama_produk'].iloc[0]}")

    ax2.set_title('ANALISIS NILAI PENJUALAN PRODUK RISING STAR\n(Nilai Penjualan Asli)', fontweight='bold', fontsize=16, pad=20)
    ax2.set_xlabel('Periode Tanggal', fontsize=12, labelpad=10)
    ax2.set_ylabel('Total Nilai Penjualan', fontsize=12, labelpad=10)
    ax2.grid(True, linestyle='--', linewidth=0.5, alpha=0.5)
    plt.xticks(rotation=45, ha='right')
    ax2.legend(final_handles, final_labels, title="Kategori Produk", bbox_to_anchor=(1.02, 1), loc='upper left')

    plt.tight_layout()
    plt.savefig('rising_star_actual.png', bbox_inches='tight')
    
    print("\n✅ SELESAI! Seluruh file berhasil dibuat di folder proyek Anda.")
else:
    print("\n⚠️ Data diproses, tetapi tidak ada produk yang memenuhi syarat Rising Star (naik 12 hari beruntun).")